# -*- coding: utf-8 -*-
##########################################################################
#
#    Rishvi Pvt Ltd
#
#    Copyright (C) 2025-TODAY Rishvi Pvt Ltd (<https://www.rishvi.com>)
#    Author: Rishvi Development Team (<https://www.rishvi.com>)
#
#    This software is proprietary and confidential.
#
#    Unauthorized copying, modification, distribution, or use of this
#    software, via any medium, is strictly prohibited without the prior
#    written permission of Rishvi Pvt Ltd.
#
#    This software is licensed, not sold. A valid commercial license
#    from Rishvi Pvt Ltd is required to use this software.
#
#    THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND,
#    EXPRESS OR IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES
#    OF MERCHANTABILITY, FITNESS FOR A PARTICULAR PURPOSE AND
#    NON-INFRINGEMENT.
#
#    For licensing information, contact: support@rishvi.com
#
##########################################################################
import base64
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError


class HrAttendanceManagerMonthlyReport(models.TransientModel):
    _name = 'hr.attendance.manager.monthly.report'
    _description = 'Attendance Manager Monthly Report'
    _rec_name="manager_user_id"

    manager_user_id = fields.Many2one('res.users', string="User", readonly=True)
    date_from = fields.Date(string="Date From", readonly=True)
    date_to = fields.Date(string="Date To", readonly=True)
    month_label = fields.Char(string="Month", readonly=True)
    managed_employee_count = fields.Integer(string="Employees in Preview", readonly=True)
    attendance_record_count = fields.Integer(string="Attendance Records", readonly=True)
    export_file = fields.Binary(string="Export File", readonly=True, attachment=False)
    export_filename = fields.Char(string="Filename", readonly=True)
    summary_line_ids = fields.One2many(
        'hr.attendance.manager.monthly.report.line', 'report_id', string="Summary", readonly=True
    )
    detail_line_ids = fields.One2many(
        'hr.attendance.manager.monthly.report.detail', 'report_id', string="Details", readonly=True
    )

    @api.model
    def action_open_report(self):
        report = self.create({})
        report._refresh_report_data()
        return report._get_open_action()

    @api.model
    def action_open_filter_report(self):
        report = self.create({})
        report._refresh_report_data()
        return report.action_open_detail_list()

    def _get_open_action(self):
        self.ensure_one()
        form_view = self.env.ref('rishvi_hrms_module_v19.view_attendance_manager_monthly_report_form')
        return {
            'type': 'ir.actions.act_window',
            'name': _('Preview'),
            'res_model': 'hr.attendance.manager.monthly.report',
            'view_mode': 'form',
            'views': [(form_view.id, 'form')],
            'target': 'current',
            'res_id': self.id,
        }

    def action_refresh(self):
        self.ensure_one()
        self._refresh_report_data()
        return self._get_open_action()

    def action_open_detail_list(self):
        self.ensure_one()
        list_view = self.env.ref(
            'rishvi_hrms_module_v19.view_attendance_manager_monthly_report_detail_list'
        )
        form_view = self.env.ref(
            'rishvi_hrms_module_v19.view_attendance_manager_monthly_report_detail_form'
        )
        search_view = self.env.ref(
            'rishvi_hrms_module_v19.view_attendance_manager_monthly_report_detail_search'
        )
        return {
            'type': 'ir.actions.act_window',
            'name': _('Attendance Detail'),
            'res_model': 'hr.attendance.manager.monthly.report.detail',
            'view_mode': 'list,form',
            'views': [(list_view.id, 'list'), (form_view.id, 'form')],
            'search_view_id': search_view.id,
            'target': 'current',
            'domain': [('report_id', '=', self.id)],
            'context': {
                'attendance_preview_report_id': self.id,
                'active_model': 'hr.attendance.manager.monthly.report',
                'active_id': self.id,
                'active_ids': [self.id],
                'search_default_group_employee': 1,
            },
        }

    def action_export_pdf(self):
        self.ensure_one()
        return self.env.ref('rishvi_hrms_module_v19.action_report_attendance_manager_monthly').report_action(self)

    def action_export_xlsx(self):
        self.ensure_one()
        file_bytes = self._generate_xlsx_bytes()
        self.write({
            'export_file': base64.b64encode(file_bytes),
            'export_filename': self._build_export_filename('xlsx'),
        })
        return self._download_export()

    def action_export_docx(self):
        self.ensure_one()
        file_bytes = self._generate_docx_bytes()
        self.write({
            'export_file': base64.b64encode(file_bytes),
            'export_filename': self._build_export_filename('docx'),
        })
        return self._download_export()

    def action_send_to_mail(self):
        self.ensure_one()
        recipient_email = self.manager_user_id.partner_id.email or self.manager_user_id.email
        if not recipient_email:
            raise UserError(_("Please add an email address on the logged-in user before sending the report mail."))
        job = self.env['hr.attendance.manager.monthly.mail.job'].sudo().create({
            'report_id': self.id,
            'requested_by_user_id': self.env.user.id,
            'recipient_email': recipient_email,
        })
        try:
            self.env.ref('rishvi_hrms_module_v19.ir_cron_attendance_report_mail_queue').sudo()._trigger(
                at=fields.Datetime.now()
            )
        except Exception:
            pass
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Mail Queued'),
                'message': _(
                    'Attendance report mail has been queued for %s. It will be sent in the background.'
                ) % recipient_email,
                'type': 'success',
                'sticky': False,
            },
        }

    def _download_export(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': (
                '/web/content?model=hr.attendance.manager.monthly.report'
                f'&id={self.id}&field=export_file&filename_field=export_filename&download=true'
            ),
            'target': 'self',
        }

    def _build_export_filename(self, extension):
        self.ensure_one()
        month_token = (self.month_label or 'current_month').replace(' ', '_').lower()
        return f'attendance_manager_current_month_{month_token}.{extension}'

    def _generate_pdf_bytes(self):
        self.ensure_one()
        report_action = self.env.ref('rishvi_hrms_module_v19.action_report_attendance_manager_monthly')
        try:
            pdf_content, report_type = report_action._render_qweb_pdf(
                report_action.report_name,
                res_ids=self.ids,
            )
        except Exception as exc:
            raise UserError(_("PDF generation failed. Please make sure wkhtmltopdf is installed.")) from exc
        if not pdf_content:
            raise UserError(_("PDF generation returned an empty file."))
        return pdf_content

    def _get_current_month_range(self):
        today = fields.Date.context_today(self)
        return today.replace(day=1), today

    def _get_scoped_public_employees(self):
        user_employee = self.env.user.employee_id
        if not user_employee:
            return self.env['hr.employee.public']
        return self.env['hr.employee.public'].search([('id', 'child_of', user_employee.id)])

    def _get_attendance_domain(self, date_from, date_to):
        start_dt = datetime.combine(date_from, time.min)
        end_dt = datetime.combine(date_to + timedelta(days=1), time.min)
        scoped_employees = self._get_scoped_public_employees()
        return [
            ('check_in', '>=', start_dt),
            ('check_in', '<', end_dt),
            ('employee_id', 'in', scoped_employees.ids),
        ]

    def _refresh_report_data(self):
        self.ensure_one()
        date_from, date_to = self._get_current_month_range()
        scoped_public_employees = self._get_scoped_public_employees()
        attendances = self.env['hr.attendance'].with_context(active_test=False).search(
            self._get_attendance_domain(date_from, date_to),
            order='employee_id, check_in',
        )
        public_employee_map = {
            public_employee.id: public_employee
            for public_employee in scoped_public_employees.exists()
        }
        employee_ids = sorted(
            {
                attendance.employee_id.id
                for attendance in attendances
                if attendance.employee_id.id in public_employee_map
            },
            key=lambda employee_id: (public_employee_map[employee_id].name or '').lower(),
        )

        attendance_map = defaultdict(list)
        for attendance in attendances:
            attendance_map[attendance.employee_id.id].append(attendance)

        summary_commands = [Command.clear()]
        detail_commands = [Command.clear()]
        summary_count = 0

        for employee_id in employee_ids:
            employee_attendances = attendance_map.get(employee_id, [])
            public_employee = public_employee_map.get(employee_id)
            if not public_employee:
                continue
            badge_id = public_employee.barcode or ''

            summary_count += 1
            present_days = len({
                attendance.check_in.date()
                for attendance in employee_attendances
                if attendance.check_in
            })
            late_count = sum(1 for attendance in employee_attendances if attendance.attendance_status == 'late')
            early_leave_count = sum(
                1 for attendance in employee_attendances if attendance.attendance_status == 'early'
            )
            total_worked_hours = sum(
                attendance.custom_worked_hours or attendance.worked_hours or 0.0
                for attendance in employee_attendances
            )
            total_extra_hours = sum(attendance.custom_extra_hours or 0.0 for attendance in employee_attendances)

            summary_commands.append(Command.create({
                'employee_id': employee_id,
                'employee_name': public_employee.name,
                'barcode': badge_id,
                'present_days': present_days,
                'late_count': late_count,
                'early_leave_count': early_leave_count,
                'total_worked_hours': total_worked_hours,
                'total_extra_hours': total_extra_hours,
            }))

            for attendance in employee_attendances:
                detail_commands.append(Command.create({
                    'employee_id': employee_id,
                    'employee_name': public_employee.name,
                    'barcode': badge_id,
                    'attendance_id': attendance.id,
                    'attendance_date': attendance.check_in.date() if attendance.check_in else False,
                    'check_in': attendance.check_in,
                    'check_out': attendance.check_out,
                    'attendance_status': dict(attendance._fields['attendance_status'].selection).get(
                        attendance.attendance_status, _('Present')
                    ),
                    'attendance_shift': attendance.attendance_shift or '',
                    'worked_hours': attendance.custom_worked_hours or attendance.worked_hours or 0.0,
                    'extra_hours': attendance.custom_extra_hours or 0.0,
                    'late_minutes': attendance.late_minutes or 0.0,
                    'early_minutes': attendance.early_minutes or 0.0,
                }))

        self.write({
            'manager_user_id': self.env.user.id,
            'date_from': date_from,
            'date_to': date_to,
            'month_label': date_from.strftime('%B %Y'),
            'managed_employee_count': summary_count,
            'attendance_record_count': len(attendances),
            'summary_line_ids': summary_commands,
            'detail_line_ids': detail_commands,
        })

    def _get_sorted_summary_lines(self):
        self.ensure_one()
        return self.summary_line_ids.sorted(key=lambda line: ((line.employee_name or '').lower(), line.id))

    def _get_sorted_detail_lines(self):
        self.ensure_one()
        return self.detail_line_ids.sorted(
            key=lambda line: (
                (line.employee_name or '').lower(),
                line.attendance_date or date.min,
                line.check_in or datetime.min,
            )
        )

    def get_export_sections(self):
        self.ensure_one()
        detail_map = defaultdict(list)
        for detail in self._get_sorted_detail_lines():
            detail_map[detail.employee_id.id].append(detail)

        sections = []
        for summary in self._get_sorted_summary_lines():
            sections.append({
                'summary': summary,
                'details': detail_map.get(summary.employee_id.id, []),
            })
        return sections

    def format_float_hours(self, value):
        total_minutes = int(round((value or 0.0) * 60))
        sign = '-' if total_minutes < 0 else ''
        total_minutes = abs(total_minutes)
        hours, minutes = divmod(total_minutes, 60)
        return f'{sign}{hours:02d}:{minutes:02d}'

    def format_datetime_for_report(self, value):
        if not value:
            return '-'
        localized = fields.Datetime.context_timestamp(self, value)
        return localized.strftime('%Y-%m-%d %H:%M')

    def _generate_xlsx_bytes(self):
        self.ensure_one()
        try:
            import xlsxwriter  # noqa: PLC0415
        except ImportError as exc:
            raise UserError(_("xlsxwriter is required in the Odoo environment for XLSX export.")) from exc

        output = BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        header_format = workbook.add_format({'bold': True, 'bg_color': '#D9E2F3', 'border': 1})
        cell_format = workbook.add_format({'border': 1})

        summary_sheet = workbook.add_worksheet('Summary')
        details_sheet = workbook.add_worksheet('Details')

        summary_sheet.write(0, 0, 'Attendance Manager Monthly Report')
        summary_sheet.write(1, 0, 'Manager')
        summary_sheet.write(1, 1, self.manager_user_id.name or '')
        summary_sheet.write(2, 0, 'Period')
        summary_sheet.write(2, 1, f'{self.date_from} to {self.date_to}')
        summary_headers = [
            'Employee',
            'Badge ID',
            'Present Days',
            'Late Count',
            'Early Leave Count',
            'Worked Hours',
            'Extra Hours',
        ]
        for col, header in enumerate(summary_headers):
            summary_sheet.write(4, col, header, header_format)

        row = 5
        for section in self.get_export_sections():
            summary = section['summary']
            summary_sheet.write(row, 0, summary.employee_name or '', cell_format)
            summary_sheet.write(row, 1, summary.barcode or '', cell_format)
            summary_sheet.write(row, 2, summary.present_days, cell_format)
            summary_sheet.write(row, 3, summary.late_count, cell_format)
            summary_sheet.write(row, 4, summary.early_leave_count, cell_format)
            summary_sheet.write(row, 5, self.format_float_hours(summary.total_worked_hours), cell_format)
            summary_sheet.write(row, 6, self.format_float_hours(summary.total_extra_hours), cell_format)
            row += 1

        details_headers = [
            'Employee',
            'Badge ID',
            'Attendance Date',
            'Check In',
            'Check Out',
            'Status',
            'Shift',
            'Worked Hours',
            'Extra Hours',
            'Late Minutes',
            'Early Leave Minutes',
        ]
        for col, header in enumerate(details_headers):
            details_sheet.write(0, col, header, header_format)

        row = 1
        for section in self.get_export_sections():
            for detail in section['details']:
                details_sheet.write(row, 0, detail.employee_name or '', cell_format)
                details_sheet.write(row, 1, detail.barcode or '', cell_format)
                details_sheet.write(row, 2, str(detail.attendance_date or ''), cell_format)
                details_sheet.write(row, 3, self.format_datetime_for_report(detail.check_in), cell_format)
                details_sheet.write(row, 4, self.format_datetime_for_report(detail.check_out), cell_format)
                details_sheet.write(row, 5, detail.attendance_status or '', cell_format)
                details_sheet.write(row, 6, detail.attendance_shift or '', cell_format)
                details_sheet.write(row, 7, self.format_float_hours(detail.worked_hours), cell_format)
                details_sheet.write(row, 8, self.format_float_hours(detail.extra_hours), cell_format)
                details_sheet.write(row, 9, detail.late_minutes or 0.0, cell_format)
                details_sheet.write(row, 10, detail.early_minutes or 0.0, cell_format)
                row += 1

        workbook.close()
        return output.getvalue()

    def _generate_docx_bytes(self):
        self.ensure_one()
        try:
            from docx import Document  # noqa: PLC0415
        except ImportError as exc:
            raise UserError(_("python-docx is required in the Odoo environment for DOCX export.")) from exc

        document = Document()
        document.add_heading('Attendance Manager Monthly Report', level=0)
        document.add_paragraph(f'Manager: {self.manager_user_id.name or ""}')
        document.add_paragraph(f'Period: {self.date_from} to {self.date_to}')

        document.add_heading('Summary', level=1)
        summary_table = document.add_table(rows=1, cols=7)
        summary_headers = [
            'Employee', 'Badge ID', 'Present Days', 'Late Count',
            'Early Leave Count', 'Worked Hours', 'Extra Hours',
        ]
        for index, header in enumerate(summary_headers):
            summary_table.rows[0].cells[index].text = header

        for section in self.get_export_sections():
            summary = section['summary']
            row = summary_table.add_row().cells
            row[0].text = summary.employee_name or ''
            row[1].text = summary.barcode or ''
            row[2].text = str(summary.present_days)
            row[3].text = str(summary.late_count)
            row[4].text = str(summary.early_leave_count)
            row[5].text = self.format_float_hours(summary.total_worked_hours)
            row[6].text = self.format_float_hours(summary.total_extra_hours)

        document.add_heading('Daily Detail', level=1)
        for section in self.get_export_sections():
            summary = section['summary']
            document.add_heading(summary.employee_name or '-', level=2)
            document.add_paragraph(f'Badge ID: {summary.barcode or "-"}')

            detail_table = document.add_table(rows=1, cols=9)
            detail_headers = [
                'Date', 'Check In', 'Check Out', 'Status', 'Shift',
                'Worked Hours', 'Extra Hours', 'Late Minutes', 'Early Leave Minutes',
            ]
            for index, header in enumerate(detail_headers):
                detail_table.rows[0].cells[index].text = header

            for detail in section['details']:
                row = detail_table.add_row().cells
                row[0].text = str(detail.attendance_date or '')
                row[1].text = self.format_datetime_for_report(detail.check_in)
                row[2].text = self.format_datetime_for_report(detail.check_out)
                row[3].text = detail.attendance_status or ''
                row[4].text = detail.attendance_shift or ''
                row[5].text = self.format_float_hours(detail.worked_hours)
                row[6].text = self.format_float_hours(detail.extra_hours)
                row[7].text = str(detail.late_minutes or 0.0)
                row[8].text = str(detail.early_minutes or 0.0)

        output = BytesIO()
        document.save(output)
        return output.getvalue()


class HrAttendanceManagerMonthlyReportLine(models.TransientModel):
    _name = 'hr.attendance.manager.monthly.report.line'
    _description = 'Attendance Manager Monthly Report Summary Line'
    _order = 'employee_name, id'

    report_id = fields.Many2one(
        'hr.attendance.manager.monthly.report', string="Report", required=True, ondelete='cascade'
    )
    employee_id = fields.Many2one('hr.employee.public', string="Employee", readonly=True)
    employee_name = fields.Char(string="Employee Name", readonly=True)
    barcode = fields.Char(string="Badge ID", readonly=True)
    present_days = fields.Integer(string="Present Days", readonly=True)
    late_count = fields.Integer(string="Late Count", readonly=True)
    early_leave_count = fields.Integer(string="Early Leave Count", readonly=True)
    total_worked_hours = fields.Float(string="Worked Hours", readonly=True)
    total_extra_hours = fields.Float(string="Extra Hours", readonly=True)


class HrAttendanceManagerMonthlyReportDetail(models.TransientModel):
    _name = 'hr.attendance.manager.monthly.report.detail'
    _description = 'Attendance Manager Monthly Report Detail Line'
    _order = 'employee_name, attendance_date, check_in, id'
    _rec_name="employee_name"
    
    report_id = fields.Many2one(
        'hr.attendance.manager.monthly.report', string="Report", required=True, ondelete='cascade'
    )
    employee_id = fields.Many2one('hr.employee.public', string="Employee", readonly=True)
    employee_name = fields.Char(string="Employee Name", readonly=True)
    barcode = fields.Char(string="Badge ID", readonly=True)
    attendance_id = fields.Many2one('hr.attendance', string="Attendance", readonly=True)
    attendance_date = fields.Date(string="Attendance Date", readonly=True)
    check_in = fields.Datetime(string="Check In", readonly=True)
    check_out = fields.Datetime(string="Check Out", readonly=True)
    attendance_status = fields.Char(string="Status", readonly=True)
    attendance_shift = fields.Char(string="Shift", readonly=True)
    worked_hours = fields.Float(string="Worked Hours", readonly=True)
    extra_hours = fields.Float(string="Extra Hours", readonly=True)
    late_minutes = fields.Float(string="Late Minutes", readonly=True)
    early_minutes = fields.Float(string="Early Leave Minutes", readonly=True)

    def _get_report_for_list_actions(self):
        report_id = self.env.context.get('attendance_preview_report_id')
        if report_id:
            report = self.env['hr.attendance.manager.monthly.report'].browse(report_id).exists()
            if report:
                return report
        if self:
            report = self[0].report_id.exists()
            if report:
                return report
        raise UserError(_("The attendance preview report could not be found. Please reopen Preview."))

    def action_export_pdf(self):
        return self._get_report_for_list_actions().action_export_pdf()

    def action_export_xlsx(self):
        return self._get_report_for_list_actions().action_export_xlsx()

    def action_export_docx(self):
        return self._get_report_for_list_actions().action_export_docx()

    def action_send_to_mail(self):
        return self._get_report_for_list_actions().action_send_to_mail()
